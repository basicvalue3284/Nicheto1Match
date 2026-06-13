"""PC batch runner for NicheTo1Match.

This is intentionally separate from the FastAPI website. It reuses the same
pipeline modules, but writes local files and a final ZIP for bulk desktop runs.
"""
from __future__ import annotations

import argparse
import asyncio
import json
import os
import re
import shutil
import sys
import uuid
from dataclasses import asdict
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

os.environ.setdefault("DATA_ROOT", str(ROOT / "Nicheto1Match_AutoRUN" / ".batch_data"))
load_dotenv(ROOT / ".env")

from openpyxl import Workbook

from app.pipeline.defaults import CLEANING_EXAMPLES, POLICY_EXAMPLES, POLICY_PROMPT
from app.pipeline.parser import parse_upload
from app.pipeline.processor import process_job
from app.pipeline.segregate import TYPE_NON_HOW_TO, segregate, summarize
from app.pipeline.title_scraper import run_bulk_scrape
from app.store.jobs import Job, append_log, load_job, make_rows, refresh_counts, save_job, utc_now
from app.store.scrapes import ScrapeJob, append_scrape_log, load_scrape_job, save_scrape_job

SUPPORTED_EXTENSIONS = {".xlsx", ".xls", ".csv", ".tsv", ".txt"}
DEFAULT_INPUT_DIR = ROOT / "Nicheto1Match_AutoRUN" / "input_titles"
DEFAULT_OUTPUT_DIR = ROOT / "Nicheto1Match_AutoRUN" / "output_results"


def main() -> None:
    args = parse_args()
    input_dir = Path(args.input_dir).expanduser().resolve()
    output_dir = Path(args.output_dir).expanduser().resolve()
    files = discover_files(input_dir)
    if args.limit_files:
        files = files[: args.limit_files]
    if not files:
        print(f"No title files found in {input_dir}")
        return

    output_dir.mkdir(parents=True, exist_ok=True)
    run_dir = output_dir / f"batch_{utc_now().replace(':', '').replace('-', '').replace('Z', '')}"
    run_dir.mkdir(parents=True, exist_ok=True)

    print(f"Found {len(files)} file(s).")
    print(f"Output: {run_dir}")
    summaries = []
    for number, path in enumerate(files, start=1):
        print(f"\n[{number}/{len(files)}] Processing {path.name}")
        try:
            summary = asyncio.run(process_file(path, run_dir, args))
            summaries.append(summary)
            print(f"  Done: {summary['safe_titles']} safe, {summary['one_match_titles']} one-match.")
        except Exception as exc:  # noqa: BLE001 - keep batch moving
            error_dir = run_dir / safe_name(path.stem)
            error_dir.mkdir(parents=True, exist_ok=True)
            summary = {
                "file": path.name,
                "status": "ERROR",
                "error": str(exc),
            }
            summaries.append(summary)
            (error_dir / "error.txt").write_text(str(exc), encoding="utf-8")
            print(f"  Error: {exc}")

    summary_path = run_dir / "batch_summary.json"
    summary_path.write_text(json.dumps(summaries, ensure_ascii=False, indent=2), encoding="utf-8")
    zip_path = make_zip(run_dir)
    print(f"\nBatch complete.")
    print(f"Summary: {summary_path}")
    print(f"ZIP: {zip_path}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run NicheTo1Match on local title files.")
    parser.add_argument("--input-dir", default=str(DEFAULT_INPUT_DIR), help="Folder with xlsx/csv/txt title files.")
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR), help="Folder where results are written.")
    parser.add_argument("--provider", default=os.getenv("AI_PROVIDER", "openai"), choices=["openai", "gemini"])
    parser.add_argument("--model", default=os.getenv("OPENAI_MODEL", "gpt-4o-mini"))
    parser.add_argument("--gemini-model", default=os.getenv("GEMINI_MODEL", "gemini-1.5-flash"))
    parser.add_argument("--batch-size", type=int, default=int(os.getenv("BATCH_SIZE", "25")))
    parser.add_argument("--skip-step3", action="store_true", help="Only run segregation + extraction/policy.")
    parser.add_argument("--dry-run-step1", action="store_true", help="Validate parsing/segregation/exports without API calls.")
    parser.add_argument("--limit-files", type=int, default=0, help="Optional test limit.")
    return parser.parse_args()


def discover_files(input_dir: Path) -> list[Path]:
    files = [
        path
        for path in sorted(input_dir.iterdir())
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS and not path.name.startswith("~$")
    ]
    return files


def parse_title_file(path: Path) -> list[str]:
    content = path.read_bytes()
    try:
        return parse_upload(path.name, content)
    except Exception:
        if path.suffix.lower() == ".xls":
            return parse_excel_xml_first_column(content)
        raise


def parse_excel_xml_first_column(content: bytes) -> list[str]:
    text = content.decode("utf-8-sig", errors="ignore").strip()
    if not text:
        return []
    if text.startswith("<?xml") or "<Workbook" in text[:500]:
        root = ElementTree.fromstring(text)
        values: list[str] = []
        for row in root.iter():
            if not row.tag.endswith("Row"):
                continue
            cells = [node for node in row if node.tag.endswith("Cell")]
            if not cells:
                continue
            first = ""
            for child in cells[0].iter():
                if child.tag.endswith("Data") and child.text:
                    first = child.text.strip()
                    break
            if first:
                values.append(first)
        if values and values[0].strip().lower() in {"title", "titles", "input title", "input titles"}:
            values = values[1:]
        return dedupe(values)
    return dedupe(line.strip().split("\t")[0].strip() for line in text.splitlines())


def dedupe(values) -> list[str]:
    seen = set()
    output: list[str] = []
    for value in values:
        title = str(value or "").strip()
        if not title:
            continue
        key = title.lower()
        if key in seen:
            continue
        seen.add(key)
        output.append(title)
    return output


async def process_file(path: Path, run_dir: Path, args: argparse.Namespace) -> dict[str, object]:
    output_dir = run_dir / safe_name(path.stem)
    if output_dir.exists():
        shutil.rmtree(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    titles = parse_title_file(path)
    if not titles:
        raise RuntimeError("No titles found in Column A.")

    job = create_phase1_job(path, titles, args)
    print(f"  Step 1: {job.raw_count} uploaded, {job.final_count} final-list titles.")
    if args.dry_run_step1:
        export_phase1(job, output_dir)
        export_final_master(job, None, output_dir)
        summary = {
            "file": path.name,
            "status": "DRY_RUN_STEP1",
            "phase1_job_id": job.id,
            "scrape_job_id": "",
            "uploaded_titles": job.raw_count,
            "unique_titles": job.total,
            "final_titles": job.final_count,
            "safe_titles": 0,
            "review_rows": job.final_count,
            "one_match_titles": 0,
            "output_folder": str(output_dir),
        }
        (output_dir / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
        return summary
    await process_job(job.id, None, reprocess_all=True)
    job = load_job(job.id)
    if not job:
        raise RuntimeError("Step 2 job disappeared from local store.")
    refresh_counts(job)
    save_job(job)

    safe_titles = [row.cleaned_title for row in job.rows if row.include_final and row.decision == "DO" and row.cleaned_title]
    review_rows = [
        row for row in job.rows
        if row.include_final and (row.decision != "DO" or row.status != "SAFE" or not row.cleaned_title)
    ]
    print(f"  Step 2: {len(safe_titles)} safe, {len(review_rows)} review/error rows.")

    export_phase1(job, output_dir)

    scrape = None
    one_match_titles: list[str] = []
    if safe_titles and not args.skip_step3:
        scrape = create_scrape_job(job, safe_titles)
        print(f"  Step 3: scraping {len(safe_titles)} safe titles...")
        await run_bulk_scrape(scrape.id, None)
        scrape = load_scrape_job(scrape.id)
        if not scrape:
            raise RuntimeError("Step 3 job disappeared from local store.")
        one_match_titles = sorted({row.input_title for row in scrape.results if row.match_count == 1 and row.matched_result})
        export_scrape(scrape, output_dir)
    elif not safe_titles:
        (output_dir / "step3_skipped.txt").write_text("Step 3 skipped: zero SAFE titles from Step 2.", encoding="utf-8")
    export_final_master(job, scrape, output_dir)

    summary = {
        "file": path.name,
        "status": "DONE",
        "phase1_job_id": job.id,
        "scrape_job_id": scrape.id if scrape else "",
        "uploaded_titles": job.raw_count,
        "unique_titles": job.total,
        "final_titles": job.final_count,
        "safe_titles": len(safe_titles),
        "review_rows": len(review_rows),
        "one_match_titles": len(one_match_titles),
        "output_folder": str(output_dir),
    }
    (output_dir / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary


def create_phase1_job(path: Path, titles: list[str], args: argparse.Namespace) -> Job:
    results = segregate(titles)
    provider = (args.provider or "openai").strip().lower()
    model = args.gemini_model if provider == "gemini" else args.model
    job = Job(
        id=uuid.uuid4().hex[:10],
        niche=path.stem.replace("_", " ").replace("-", " ").strip() or "Batch titles",
        created_at=utc_now(),
        updated_at=utc_now(),
        state="QUEUED",
        total=len(results),
        processed=0,
        failed=0,
        batch_size=max(5, min(int(args.batch_size or 25), 50)),
        model=model,
        provider=provider,
        summary=summarize(results),
        cleaning_examples=CLEANING_EXAMPLES,
        policy_examples=POLICY_EXAMPLES,
        policy_prompt=POLICY_PROMPT,
        rows=make_rows(results),
        raw_count=len(titles),
        duplicate_count=0,
        auto_export=True,
    )
    append_log(job, f"PC batch created from {path.name}.")
    save_job(job)
    return job


def create_scrape_job(job: Job, titles: list[str]) -> ScrapeJob:
    scrape = ScrapeJob(
        id=uuid.uuid4().hex[:10],
        name=f"{job.niche} - Title Scraper",
        created_at=utc_now(),
        updated_at=utc_now(),
        state="QUEUED",
        titles=titles,
        source_phase1_job_id=job.id,
    )
    append_scrape_log(scrape, f"PC batch created from Phase 1 job {job.id} with {len(titles)} title(s).")
    save_scrape_job(scrape)
    return scrape


def export_phase1(job: Job, output_dir: Path) -> None:
    headers = ["#", "Original Title", "Title Extraction", "Removed Text", "Tag", "Policy", "Reason Notes", "Confidence", "Error"]
    groups = [
        ("Raw Data", job.rows),
        ("How-To", [row for row in job.rows if row.type != TYPE_NON_HOW_TO]),
        ("How-To Start", [row for row in job.rows if row.type == "How-To (Start)"]),
        ("How-To Board", [row for row in job.rows if row.type == "How-To (Board)"]),
        ("How-To Shorts", [row for row in job.rows if row.type == "How-To (Shorts)"]),
        ("Non How-To", [row for row in job.rows if row.type == TYPE_NON_HOW_TO]),
        ("Final List", [row for row in job.rows if row.include_final]),
        ("Review Rows", [
            row for row in job.rows
            if row.include_final and (row.decision != "DO" or row.status != "SAFE" or not row.cleaned_title)
        ]),
    ]
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, rows in groups:
        ws = workbook.create_sheet(sheet_name[:31])
        ws.append(headers)
        for idx, row in enumerate(rows, start=1):
            ws.append(phase1_row(row, idx))
        autosize(ws)
    workbook.save(output_dir / "phase1_master.xlsx")

    write_single_column_csv(output_dir / "safe_titles.csv", "Title", [
        row.cleaned_title for row in job.rows if row.include_final and row.decision == "DO" and row.cleaned_title
    ])
    review_rows = [
        row for row in job.rows
        if row.include_final and (row.decision != "DO" or row.status != "SAFE" or not row.cleaned_title)
    ]
    review_book = Workbook()
    ws = review_book.active
    ws.title = "Review Rows"
    ws.append(headers)
    for idx, row in enumerate(review_rows, start=1):
        ws.append(phase1_row(row, idx))
    autosize(ws)
    review_book.save(output_dir / "review_rows.xlsx")


def phase1_row(row: object, idx: int) -> list[str]:
    return [
        str(idx),
        row.original_title,
        row.cleaned_title,
        "; ".join(row.removed_words),
        row.type,
        row.status,
        row.reason,
        row.confidence,
        row.error,
    ]


def export_scrape(job: ScrapeJob, output_dir: Path) -> None:
    headers = scrape_headers()
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Step 3 Master"
    ws.append(headers)
    for row in job.results:
        ws.append(scrape_row(row))
    autosize(ws)
    workbook.save(output_dir / "step3_master.xlsx")

    one_match = []
    seen = set()
    for row in job.results:
        if row.match_count == 1 and row.matched_result and row.input_title not in seen:
            seen.add(row.input_title)
            one_match.append(row)
    one_match = sorted(one_match, key=status_sort_key)
    write_single_column_csv(output_dir / "Final 1 Match_Titles.csv", "Title", [row.input_title for row in one_match])

    details = Workbook()
    ws = details.active
    ws.title = "Final 1 Match with Stats"
    ws.append(["Title", "Views", "Subs", "Match Count", "Status Tagging", "Channel ID", "Result Title"])
    for row in one_match:
        ws.append([row.input_title, row.views, row.subs, row.match_count, row.status_tagging, row.channel_id, row.result_title])
    autosize(ws)
    details.save(output_dir / "Final 1 Match with Stats.xlsx")


def status_sort_key(row: object) -> tuple[int, int, str]:
    status = str(row.status_tagging or "").lower()
    if "<5k" in status:
        bucket = 0
    elif "5k-10k" in status or "5k to 10k" in status:
        bucket = 1
    elif "10k-25k" in status or "10k to 25k" in status:
        bucket = 2
    elif "25k+" in status:
        bucket = 3
    else:
        bucket = 9
    return bucket, parse_subscriber_value(row.subs), str(row.input_title).lower()


def parse_subscriber_value(value: str) -> int:
    text = str(value or "").strip().lower().replace(",", "")
    match = re.search(r"(\d+(?:\.\d+)?)\s*([kmb])?", text)
    if not match:
        return 10**12
    number = float(match.group(1))
    suffix = match.group(2)
    multiplier = {"k": 1_000, "m": 1_000_000, "b": 1_000_000_000}.get(suffix, 1)
    return int(number * multiplier)


def export_final_master(phase1: Job, scrape: ScrapeJob | None, output_dir: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)

    ws = workbook.create_sheet("Step 1")
    ws.append(["#", "Title", "Tag", "Final List"])
    for idx, row in enumerate(phase1.rows, start=1):
        ws.append([idx, row.original_title, row.type, "YES" if row.include_final else ""])
    autosize(ws)

    ws = workbook.create_sheet("Step 2 Title Extractor")
    ws.append(["#", "Original Title", "Title Extraction", "Removed Text", "Tag", "Policy", "Reason Notes", "Confidence", "Error"])
    for idx, row in enumerate([row for row in phase1.rows if row.include_final], start=1):
        ws.append(phase1_row(row, idx))
    autosize(ws)

    ws = workbook.create_sheet("Step 3")
    ws.append(scrape_headers())
    if scrape:
        for row in scrape.results:
            ws.append(scrape_row(row))
    else:
        ws.append(["Step 3 not run yet", "", "", "", "", "", "", "", "", "", ""])
    autosize(ws)

    workbook.save(output_dir / "Final_Master.xlsx")


def scrape_headers() -> list[str]:
    return [
        "Input Title",
        "Result Title",
        "Result Views",
        "Channel ID",
        "Channel Title",
        "Subs",
        "Search Results Position",
        "Match Count",
        "Status Tagging",
        "Matched Result",
        "Error",
    ]


def scrape_row(row: object) -> list[str | int]:
    return [
        row.input_title,
        row.result_title,
        row.views,
        row.channel_id,
        row.channel_title,
        row.subs,
        row.search_position,
        row.match_count,
        row.status_tagging,
        "YES" if row.matched_result else "",
        row.error,
    ]


def write_single_column_csv(path: Path, header: str, values: list[str]) -> None:
    lines = [header] + values
    path.write_text("\n".join(csv_escape(value) for value in lines) + "\n", encoding="utf-8")


def csv_escape(value: str) -> str:
    text = str(value or "")
    if any(ch in text for ch in [",", '"', "\n"]):
        return '"' + text.replace('"', '""') + '"'
    return text


def autosize(ws) -> None:
    for column in ws.columns:
        letter = column[0].column_letter
        max_len = 0
        for cell in column[:2000]:
            max_len = max(max_len, len(str(cell.value or "")))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 70)
    ws.freeze_panes = "A2"


def make_zip(run_dir: Path) -> Path:
    zip_path = run_dir.with_suffix(".zip")
    if zip_path.exists():
        zip_path.unlink()
    with ZipFile(zip_path, "w", ZIP_DEFLATED) as archive:
        for path in sorted(run_dir.rglob("*")):
            if path.is_file():
                archive.write(path, path.relative_to(run_dir.parent))
    return zip_path


def safe_name(value: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", value.strip())
    return safe.strip("._-")[:80] or "titles"


if __name__ == "__main__":
    main()
