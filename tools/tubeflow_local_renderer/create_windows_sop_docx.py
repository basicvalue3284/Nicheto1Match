from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "TubeFlow_Windows_Freelancer_SOP.docx"


def add_code_block(doc, text):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.space_after = Pt(6)


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbers(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def main():
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    for style_name, size, color in [
        ("Heading 1", 16, RGBColor(31, 78, 121)),
        ("Heading 2", 12, RGBColor(31, 78, 121)),
        ("Heading 3", 10.5, RGBColor(68, 68, 68)),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color

    title = doc.add_paragraph()
    title_run = title.add_run("TubeFlow Local Renderer SOP")
    title_run.bold = True
    title_run.font.name = "Arial"
    title_run.font.size = Pt(24)
    title_run.font.color.rgb = RGBColor(31, 78, 121)

    subtitle = doc.add_paragraph()
    subtitle_run = subtitle.add_run("Windows freelancer setup, OBS settings, and daily production workflow")
    subtitle_run.font.name = "Arial"
    subtitle_run.font.size = Pt(11)
    subtitle_run.font.color.rgb = RGBColor(90, 90, 90)

    doc.add_paragraph("Purpose", style="Heading 1")
    doc.add_paragraph(
        "This setup creates the final TubeFlow video locally on the freelancer's Windows PC. "
        "The freelancer records the screen while TubeFlow audio plays at 1.5x. The local renderer "
        "slows the video back to normal speed, removes OBS audio, adds the clean TubeFlow 1x audio, "
        "and exports the final MP4 for Google Drive upload."
    )

    doc.add_paragraph("Folder Meaning", style="Heading 1")
    folder_rows = [
        ("Downloads", "TubeFlow 1x audio is downloaded here.", r"C:\Users\<name>\Downloads\topic_name.mp3"),
        ("Videos\\LosslessCut", "LosslessCut exports the trimmed OBS video here.", r"C:\Users\<name>\Videos\LosslessCut\topic_name.mp4"),
        ("Desktop\\Ready_For_Google_Drive", "Final video appears here.", r"C:\Users\<name>\Desktop\Ready_For_Google_Drive\topic_name_FINAL.mp4"),
        ("Videos\\LosslessCut\\Processed_By_TubeFlow", "Raw trimmed video is moved here after successful render.", "Prevents duplicate rendering."),
        ("Videos\\LosslessCut\\Needs_Admin_Help", "Failed videos are moved here.", "Admin should check naming, audio, or FFmpeg."),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Folder"
    hdr[1].text = "Purpose"
    hdr[2].text = "Example / Note"
    for folder, purpose, example in folder_rows:
        cells = table.add_row().cells
        cells[0].text = folder
        cells[1].text = purpose
        cells[2].text = example

    doc.add_paragraph("One-Time Setup", style="Heading 1")
    doc.add_paragraph("1. Install Python", style="Heading 2")
    add_numbers(doc, [
        "Open https://www.python.org/downloads/windows/",
        "Download Python for Windows.",
        "During installation, tick Add python.exe to PATH.",
        "Complete installation.",
    ])

    doc.add_paragraph("2. Install FFmpeg", style="Heading 2")
    doc.add_paragraph("Official FFmpeg page:")
    add_code_block(doc, "https://ffmpeg.org/download.html")
    doc.add_paragraph("Recommended Windows build:")
    add_code_block(doc, "https://www.gyan.dev/ffmpeg/builds/")
    doc.add_paragraph("Simplest Windows 10/11 install method:")
    add_numbers(doc, [
        "Press Windows key.",
        "Search Command Prompt.",
        "Right-click and choose Run as administrator.",
        "Paste: winget install Gyan.FFmpeg",
        "Press Enter.",
        "After installation, close Command Prompt and open it again.",
        "Run ffmpeg -version to confirm it is working.",
    ])
    doc.add_paragraph("Video tutorial search link:")
    add_code_block(doc, "https://www.youtube.com/results?search_query=install+ffmpeg+windows+11+add+to+path")

    doc.add_paragraph("3. Copy Renderer Folder", style="Heading 2")
    doc.add_paragraph("Copy the full renderer folder to:")
    add_code_block(doc, r"C:\TubeFlow_Local_Renderer")
    doc.add_paragraph("The folder must contain:")
    add_bullets(doc, ["renderer.py", "config.json", "Start_TubeFlow_Renderer.bat"])

    doc.add_paragraph("4. Create Working Folders", style="Heading 2")
    add_code_block(doc, r"C:\Users\<name>\Videos\LosslessCut")
    add_code_block(doc, r"C:\Users\<name>\Desktop\Ready_For_Google_Drive")

    doc.add_paragraph("OBS Settings", style="Heading 1")
    doc.add_paragraph("Video", style="Heading 2")
    add_code_block(
        doc,
        "Base Canvas Resolution: 1920x1080\n"
        "Output Scaled Resolution: 1920x1080\n"
        "Downscale Filter: Lanczos\n"
        "Common FPS Values: 30",
    )
    doc.add_paragraph("Output > Recording", style="Heading 2")
    add_code_block(
        doc,
        "Output Mode: Advanced\n"
        "Recording Format: MKV\n"
        "Video Encoder: Hardware H.264 if available\n"
        "Audio Encoder: AAC\n"
        "Rate Control: CBR\n"
        "Bitrate: 15000 Kbps\n"
        "Keyframe Interval: 2 s\n"
        "Preset: Quality\n"
        "Profile: High",
    )
    doc.add_paragraph("Audio", style="Heading 2")
    add_code_block(doc, "Sample Rate: 44.1 kHz\nChannels: Mono or Stereo")
    doc.add_paragraph(
        "The final renderer removes OBS audio and uses TubeFlow clean audio, so OBS audio is only for reference."
    )

    doc.add_paragraph("Daily Work Process", style="Heading 1")
    doc.add_paragraph("1. Start Renderer", style="Heading 2")
    doc.add_paragraph("Before creating videos, double-click:")
    add_code_block(doc, r"C:\TubeFlow_Local_Renderer\Start_TubeFlow_Renderer.bat")
    doc.add_paragraph("Keep the black window open. It should say:")
    add_code_block(doc, "Waiting for LosslessCut exports...")

    doc.add_paragraph("2. Create Audio In TubeFlow", style="Heading 2")
    add_numbers(doc, [
        "Open TubeFlow.",
        "Create the video topic.",
        "Generate script and audio.",
        "Download the 1x audio file.",
        "Keep the audio in Downloads.",
        "Rename the audio clearly.",
    ])
    doc.add_paragraph("Good filename:")
    add_code_block(doc, "how_to_enable_cloudflare_schema_validation.mp3")
    doc.add_paragraph("Avoid filenames like audio.mp3, download (4).mp3, or final audio.mp3.")

    doc.add_paragraph("3. Record Screen In OBS", style="Heading 2")
    add_numbers(doc, [
        "Open OBS.",
        "Start recording.",
        "In TubeFlow, click Start Recording or play audio.",
        "Audio plays at 1.5x.",
        "Follow the steps on screen.",
        "Stop recording after the task is complete.",
    ])

    doc.add_paragraph("4. Trim In LosslessCut", style="Heading 2")
    doc.add_paragraph("Freelancers must use LosslessCut before final render.")
    add_bullets(doc, ["Trim delay at the beginning.", "Trim delay at the end.", "Do not cut important middle parts."])
    doc.add_paragraph("Export to:")
    add_code_block(doc, r"C:\Users\<name>\Videos\LosslessCut")
    doc.add_paragraph("Name the exported video the same as the audio:")
    add_code_block(doc, "how_to_enable_cloudflare_schema_validation.mp4")

    doc.add_paragraph("5. Wait For Final Output", style="Heading 2")
    doc.add_paragraph("The renderer will show:")
    add_code_block(doc, "Found video\nMatched audio\nStarting FFmpeg render\nFinal video ready")
    doc.add_paragraph("Upload only the final file from:")
    add_code_block(doc, r"C:\Users\<name>\Desktop\Ready_For_Google_Drive")

    doc.add_paragraph("Quality Check Before Upload", style="Heading 1")
    add_bullets(doc, [
        "Screen text is sharp.",
        "Audio is clean.",
        "Audio matches the screen action.",
        "Start delay is removed.",
        "End delay is removed.",
        "Final file name ends with _FINAL.mp4.",
    ])

    doc.add_paragraph("Important Rules", style="Heading 1")
    add_numbers(doc, [
        "Keep the renderer window open during work.",
        "Audio and video names must match.",
        "Download only one audio at a time when testing.",
        "Always trim in LosslessCut before final render.",
        "Upload only the file from Ready_For_Google_Drive.",
        "Do not upload the raw OBS recording.",
        "Do not upload the file from Processed_By_TubeFlow.",
    ])

    doc.add_paragraph("Problem Handling", style="Heading 1")
    doc.add_paragraph("If the final video does not appear, check:")
    add_code_block(doc, r"C:\Users\<name>\Videos\LosslessCut\Needs_Admin_Help")
    doc.add_paragraph("If audio is wrong, stop work and contact admin. Do not continue creating many videos.")
    doc.add_paragraph("If the renderer says FFmpeg not found, FFmpeg is not installed correctly.")

    doc.add_paragraph("Admin Render Profile", style="Heading 1")
    add_code_block(
        doc,
        "Speed correction: setpts=1.5*PTS\n"
        "Resolution: 1920x1080\n"
        "FPS: 30\n"
        "Video codec: H.264 libx264\n"
        "Quality: CRF 18\n"
        "Preset: veryfast\n"
        "Audio codec: AAC\n"
        "Audio bitrate: 96k\n"
        "Audio sample rate: 44100 Hz\n"
        "Audio channels: mono",
    )

    doc.add_page_break()
    doc.add_paragraph("Quick Checklist", style="Heading 1")
    add_bullets(doc, [
        "Renderer window is open.",
        "TubeFlow 1x audio is in Downloads.",
        "OBS recording is trimmed in LosslessCut.",
        "LosslessCut export is in Videos\\LosslessCut.",
        "Audio and video names match.",
        "Final _FINAL.mp4 is uploaded to Google Drive.",
    ])

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
