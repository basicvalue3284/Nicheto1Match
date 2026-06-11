from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from create_windows_local_render_install_doc import (
    ASSET_DIR,
    make_daily_image,
    make_flow_image,
    make_folder_image,
)


ROOT = Path(__file__).resolve().parent
INSTALL_DOC = ROOT / "TubeFlow_Local_Render_WINDOWS_INSTALL_ONLY.docx"
WORKFLOW_DOC = ROOT / "TubeFlow_Local_Render_DAILY_WORKFLOW.docx"

BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 41, 55)
MUTED = RGBColor(75, 85, 99)
LIGHT_FILL = "E8EEF5"
GREEN_FILL = "EAF7EF"


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.bold = bold
    if color:
        r.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def setup_doc(title, subtitle):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.25
    for name, size, color in [
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 12, RGBColor(31, 77, 120)),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.font.name = "Calibri"
    r.font.size = Pt(25)
    r.font.bold = True
    r.font.color.rgb = BLUE

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(subtitle)
    r.font.name = "Calibri"
    r.font.size = Pt(12)
    r.font.color.rgb = MUTED
    return doc


def add_note(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade_cell(cell, GREEN_FILL)
    set_cell_text(cell, text, color=DARK)
    doc.add_paragraph()


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(9)
    r.font.color.rgb = DARK


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_numbers(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_link_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for i, h in enumerate(["Tool", "Link", "Purpose"]):
        shade_cell(table.rows[0].cells[i], LIGHT_FILL)
        set_cell_text(table.rows[0].cells[i], h, bold=True, color=DARK)
    for tool, url, purpose in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], tool, bold=True, color=DARK)
        cells[1].text = ""
        add_hyperlink(cells[1].paragraphs[0], url, url)
        set_cell_text(cells[2], purpose, color=DARK)


def add_footer(doc, text):
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(text)
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED


def create_install_doc():
    doc = setup_doc(
        "TubeFlow Local Render - Windows Install Only",
        "One-time setup guide for freelancer PC",
    )
    add_note(
        doc,
        "Use this document only once during setup. After setup is complete, use the separate Daily Workflow document for normal work.",
    )

    doc.add_paragraph("What This Installs", style="Heading 1")
    doc.add_paragraph(
        "This installs the local renderer on the freelancer's Windows computer. The renderer automatically creates the final TubeFlow video after the freelancer places the 1x audio in folder 01 and the OBS or LosslessCut video in folder 02."
    )
    doc.add_picture(str(ASSET_DIR / "setup_flow.png"), width=Inches(7.0))

    doc.add_paragraph("What You Will Receive From Admin", style="Heading 1")
    doc.add_paragraph(
        "Admin will provide one ZIP file named TubeFlow_Local_Render_Freelancer_Package.zip. Download this ZIP file first, then extract it on the Windows PC."
    )
    doc.add_paragraph("After extracting, the folder should contain these important files:")
    add_bullets(
        doc,
        [
            "01_Audio_1x - put the downloaded TubeFlow 1x audio here.",
            "02_OBS_or_LosslessCut_Video - put the final screen recording here after trimming if needed.",
            "03_Final_Output - collect the final _FINAL.mp4 video from here.",
            "04_Processed_Raw_Video - old raw video is moved here after successful render.",
            "05_Needs_Checking - failed videos are moved here for admin checking.",
            "Start_TubeFlow_Renderer.bat - double-click this to start TubeFlow Local Render.",
            "renderer.py - main renderer program.",
            "config.json - renderer settings.",
            "TubeFlow_Local_Render_WINDOWS_INSTALL_ONLY.docx - this setup document.",
            "TubeFlow_Local_Render_DAILY_WORKFLOW.docx - daily work process document.",
        ],
    )
    add_note(
        doc,
        r"Do not create your own renderer folder. Use the ZIP file provided by admin, then copy the extracted TubeFlow_Local_Render folder to C:\TubeFlow_Local_Render.",
    )

    doc.add_paragraph("Required Links", style="Heading 1")
    add_link_table(
        doc,
        [
            ("Python for Windows", "https://www.python.org/downloads/windows/", "Required to run the renderer."),
            ("FFmpeg official page", "https://ffmpeg.org/download.html", "Required for video merging."),
            ("FFmpeg Windows builds", "https://www.gyan.dev/ffmpeg/builds/", "Recommended Windows build."),
            ("FFmpeg install video search", "https://www.youtube.com/results?search_query=install+ffmpeg+windows+11+add+to+path", "Use if extra visual help is needed."),
            ("OBS Studio", "https://obsproject.com/download", "Screen recording software."),
            ("LosslessCut", "https://github.com/mifi/lossless-cut/releases", "Used to trim start/end delay."),
        ],
    )

    doc.add_paragraph("Step 1: Install Python", style="Heading 1")
    add_numbers(
        doc,
        [
            "Open the Python link above.",
            "Download Python for Windows.",
            "On the first installer screen, tick Add python.exe to PATH.",
            "Click Install Now and finish setup.",
        ],
    )
    add_note(doc, "Important: If Add python.exe to PATH is missed, the renderer may not start.")

    doc.add_paragraph("Step 2: Install FFmpeg", style="Heading 1")
    add_numbers(
        doc,
        [
            "Press the Windows key.",
            "Search Command Prompt.",
            "Right-click Command Prompt and choose Run as administrator.",
            "Paste this command and press Enter:",
        ],
    )
    add_code(doc, "winget install Gyan.FFmpeg")
    add_numbers(
        doc,
        [
            "When Windows asks for permission, accept it.",
            "After install finishes, close Command Prompt.",
            "Open Command Prompt again and run:",
        ],
    )
    add_code(doc, "ffmpeg -version")
    add_note(doc, "If version details appear, FFmpeg is working. If not, contact admin.")

    doc.add_paragraph("Step 3: Copy The Renderer Folder", style="Heading 1")
    doc.add_paragraph(
        "After extracting the ZIP file from admin, copy the full folder named TubeFlow_Local_Render to the C drive. Do not copy only one file."
    )
    add_code(doc, r"C:\TubeFlow_Local_Render")
    doc.add_paragraph("The folder must contain these files:")
    add_bullets(
        doc,
        [
            "01_Audio_1x",
            "02_OBS_or_LosslessCut_Video",
            "03_Final_Output",
            "04_Processed_Raw_Video",
            "05_Needs_Checking",
            "renderer.py",
            "config.json",
            "Start_TubeFlow_Renderer.bat",
            "TubeFlow_Local_Render_WINDOWS_INSTALL_ONLY.docx",
            "TubeFlow_Local_Render_DAILY_WORKFLOW.docx",
        ],
    )
    doc.add_paragraph("Correct final example:")
    add_code(doc, r"C:\TubeFlow_Local_Render\Start_TubeFlow_Renderer.bat")
    add_code(doc, r"C:\TubeFlow_Local_Render\renderer.py")
    add_code(doc, r"C:\TubeFlow_Local_Render\config.json")

    doc.add_paragraph("Step 4: Check The Five Work Folders", style="Heading 1")
    doc.add_paragraph("These folders are already included in the ZIP. Open the renderer folder and confirm they are visible:")
    add_code(doc, r"C:\TubeFlow_Local_Render\01_Audio_1x")
    add_code(doc, r"C:\TubeFlow_Local_Render\02_OBS_or_LosslessCut_Video")
    add_code(doc, r"C:\TubeFlow_Local_Render\03_Final_Output")
    add_code(doc, r"C:\TubeFlow_Local_Render\04_Processed_Raw_Video")
    add_code(doc, r"C:\TubeFlow_Local_Render\05_Needs_Checking")
    doc.add_picture(str(ASSET_DIR / "folder_map.png"), width=Inches(7.0))

    doc.add_paragraph("Step 5: Test Renderer Opens", style="Heading 1")
    add_numbers(
        doc,
        [
            r"Open C:\TubeFlow_Local_Render.",
            "Double-click Start_TubeFlow_Renderer.bat.",
            "A black window should open.",
            "It should say:",
        ],
    )
    add_code(doc, "Waiting for video files...")
    add_note(doc, "Setup is complete when the black renderer window opens successfully.")

    doc.add_paragraph("OBS Settings To Confirm", style="Heading 1")
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    for i, h in enumerate(["Setting", "Value"]):
        shade_cell(table.rows[0].cells[i], LIGHT_FILL)
        set_cell_text(table.rows[0].cells[i], h, bold=True, color=DARK)
    for setting, value in [
        ("Base Canvas Resolution", "1920x1080"),
        ("Output Scaled Resolution", "1920x1080"),
        ("Downscale Filter", "Lanczos"),
        ("FPS", "30"),
        ("Recording Format", "MKV preferred, MP4 allowed"),
        ("Video Encoder", "Hardware H.264 if available"),
        ("Bitrate", "15000 Kbps"),
        ("Audio Sample Rate", "44.1 kHz"),
        ("Recording Path", r"C:\Users\<your-name>\Videos\OBS_Recordings"),
    ]:
        cells = table.add_row().cells
        set_cell_text(cells[0], setting, bold=True, color=DARK)
        set_cell_text(cells[1], value, color=DARK)

    doc.add_paragraph("Opening OBS Recording In LosslessCut", style="Heading 1")
    doc.add_paragraph(
        "OBS does not reliably auto-open every finished recording in LosslessCut. Use this simple process instead:"
    )
    add_numbers(
        doc,
        [
            r"In OBS, keep Recording Path as C:\Users\<your-name>\Videos\OBS_Recordings.",
            "After stopping OBS, click File > Show Recordings.",
            "Double-click the latest recording to open it in LosslessCut.",
            r"If trimming is needed, export the trimmed file back to C:\TubeFlow_Local_Render\02_OBS_or_LosslessCut_Video.",
        ],
    )
    add_note(
        doc,
        r"Important: do not use 02_OBS_or_LosslessCut_Video as the raw OBS recording folder if you need to trim. Put only the ready video there.",
    )

    add_footer(doc, "TubeFlow Local Render - Install Only")
    doc.save(INSTALL_DOC)


def create_workflow_doc():
    doc = setup_doc(
        "TubeFlow Local Render - Daily Workflow",
        "Use this every day after Windows setup is already completed",
    )
    add_note(
        doc,
        "Use this document for daily work. Do not repeat the install steps unless admin asks you to reinstall.",
    )
    doc.add_picture(str(ASSET_DIR / "daily_process.png"), width=Inches(7.0))

    doc.add_paragraph("Before Starting Videos", style="Heading 1")
    add_numbers(
        doc,
        [
            r"Open C:\TubeFlow_Local_Render.",
            "Double-click Start_TubeFlow_Renderer.bat.",
            "Keep the black renderer window open while working.",
            "Confirm it says Waiting for video files.",
        ],
    )

    doc.add_paragraph("For Each Video", style="Heading 1")
    doc.add_paragraph("1. Create Audio In TubeFlow", style="Heading 2")
    add_numbers(
        doc,
        [
            "Open TubeFlow Production Tool.",
            "Create the topic.",
            "Generate script and audio.",
            "Click Download 1x Audio.",
            r"Move or save the audio file into C:\TubeFlow_Local_Render\01_Audio_1x.",
        ],
    )
    add_note(doc, "Do not use random names like audio.mp3 or download (4).mp3. The audio name should match the topic.")

    doc.add_paragraph("2. Record Screen In OBS", style="Heading 2")
    add_numbers(
        doc,
        [
            "Start OBS recording.",
            "In TubeFlow, click Start Recording.",
            "The audio plays at 1.5x.",
            "Record the browser steps.",
            "Stop OBS when finished.",
        ],
    )

    doc.add_paragraph("3. Trim In LosslessCut", style="Heading 2")
    add_numbers(
        doc,
        [
            "In OBS, click File > Show Recordings.",
            "Double-click the latest OBS recording to open it in LosslessCut if trimming is needed.",
            "Trim the delay at the beginning.",
            "Trim the delay at the end.",
            r"Export the trimmed file to C:\TubeFlow_Local_Render\02_OBS_or_LosslessCut_Video.",
            "Name the exported video similar to the audio/topic.",
        ],
    )
    add_note(
        doc,
        r"Folder 02 is the renderer drop folder. Put the final screen recording there only after trimming if trimming is needed.",
    )

    doc.add_paragraph("4. Wait For Final Render", style="Heading 2")
    doc.add_paragraph("The renderer window should show:")
    add_code(doc, "Found video\nMatched audio\nStarting FFmpeg render\nFinal video ready")
    add_note(doc, "Do not close the renderer while it says Starting FFmpeg render.")

    doc.add_paragraph("5. Upload Final Video", style="Heading 2")
    add_numbers(
        doc,
        [
            r"Open C:\TubeFlow_Local_Render\03_Final_Output.",
            "Play the final video once.",
            "Upload only the file ending with _FINAL.mp4 to Google Drive.",
        ],
    )

    doc.add_paragraph("Quality Check Before Upload", style="Heading 1")
    add_bullets(
        doc,
        [
            "Screen text is sharp.",
            "Audio is clean.",
            "Audio timing matches the screen action.",
            "Start delay is removed.",
            "End delay is removed.",
            "Filename ends with _FINAL.mp4.",
        ],
    )

    doc.add_paragraph("Do Not Do These", style="Heading 1")
    add_bullets(
        doc,
        [
            "Do not close the renderer while it is rendering.",
            "Do not upload raw OBS recordings.",
            "Do not upload files from 04_Processed_Raw_Video.",
            "Do not continue many videos if audio is wrong.",
            "Do not put video files in 01_Audio_1x.",
        ],
    )

    doc.add_paragraph("If Something Goes Wrong", style="Heading 1")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for i, h in enumerate(["Problem", "Likely Reason", "What To Do"]):
        shade_cell(table.rows[0].cells[i], LIGHT_FILL)
        set_cell_text(table.rows[0].cells[i], h, bold=True, color=DARK)
    for problem, reason, fix in [
        ("Audio not found", "Audio name does not match video topic.", "Check 01_Audio_1x and contact admin."),
        ("Final video not created", "Renderer moved the file for checking.", r"Check 05_Needs_Checking."),
        ("Wrong audio", "Wrong audio was downloaded or matched.", "Stop and contact admin immediately."),
        ("Renderer closed", "Black window was closed.", "Open Start_TubeFlow_Renderer.bat again."),
        ("FFmpeg not found", "Install was not completed correctly.", "Contact admin."),
    ]:
        cells = table.add_row().cells
        set_cell_text(cells[0], problem, bold=True, color=DARK)
        set_cell_text(cells[1], reason, color=DARK)
        set_cell_text(cells[2], fix, color=DARK)

    add_footer(doc, "TubeFlow Local Render - Daily Workflow")
    doc.save(WORKFLOW_DOC)


def main():
    ASSET_DIR.mkdir(exist_ok=True)
    make_flow_image(ASSET_DIR / "setup_flow.png")
    make_folder_image(ASSET_DIR / "folder_map.png")
    make_daily_image(ASSET_DIR / "daily_process.png")
    create_install_doc()
    create_workflow_doc()
    print(INSTALL_DOC)
    print(WORKFLOW_DOC)


if __name__ == "__main__":
    main()
