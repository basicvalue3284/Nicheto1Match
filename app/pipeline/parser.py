"""Parse uploaded files or pasted text into a normalized list of titles."""
from __future__ import annotations
import csv
import io
from typing import List
from openpyxl import load_workbook


def parse_text(text: str) -> List[str]:
    """One title per line. Strips whitespace, drops empty lines and obvious duplicates."""
    seen = set()
    out: List[str] = []
    for raw in text.splitlines():
        t = raw.strip()
        if not t:
            continue
        key = t.lower()
        if key in seen:
            continue
        seen.add(key)
        out.append(t)
    return out


def parse_csv(content: bytes) -> List[str]:
    """Read CSV bytes; take first column unless a column is named 'title' (case-insensitive)."""
    text = content.decode("utf-8-sig", errors="ignore")
    sample = text[:2048]
    delimiter = "\t" if "\t" in sample and sample.count("\t") >= sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = [row for row in reader if row]
    return _select_title_column(rows)


def parse_xlsx(content: bytes) -> List[str]:
    """Read XLSX bytes; take first column unless a 'title' column exists."""
    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    ws = wb.worksheets[0]
    rows: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        values = ["" if cell is None else str(cell).strip() for cell in row]
        if any(values):
            rows.append(values)
    return _select_title_column(rows)


def _select_title_column(rows: list[list[str]]) -> List[str]:
    if not rows:
        return []
    header = rows[0]
    title_idx = None
    for idx, col in enumerate(header):
        name = col.strip().lower()
        if name in ("title", "raw title", "input titles", "input title", "titles"):
            title_idx = idx
            break
    data_rows = rows[1:] if title_idx is not None else rows
    if title_idx is None:
        title_idx = 0
    vals = [row[title_idx].strip() for row in data_rows if len(row) > title_idx]
    return parse_text("\n".join(vals))


def parse_upload(filename: str, content: bytes) -> List[str]:
    """Dispatch on file extension."""
    name = (filename or "").lower()
    if name.endswith(".csv") or name.endswith(".tsv") or name.endswith(".txt"):
        return parse_csv(content)
    if name.endswith(".xlsx") or name.endswith(".xls"):
        return parse_xlsx(content)
    # Fall back: treat as plain text
    return parse_text(content.decode("utf-8", errors="ignore"))
